import io
from datetime import date

from django.contrib.auth.models import User
from django.test import RequestFactory, TestCase
from openpyxl import load_workbook

from mi_app.models import (
    Cliente, Permiso, Prestamo, Rol, RolPermiso, UsuarioProfile,
)
from mi_app.views_core import exportar_clientes_excel, exportar_prestamos_excel


class ExcelFormulaInjectionTests(TestCase):
    """
    Regresión de auditoría de seguridad (Formula/CSV Injection, CWE-1236):
    un nombre de cliente que empiece con '=', '+', '-' o '@' se escribia
    crudo en las celdas de las exportaciones a Excel. Excel lo interpreta
    como formula al abrir el archivo, no como texto -- quien lo escribe
    (cualquier operario) no es quien despues lo abre (tipicamente un
    ADMIN/GERENTE generando reportes). Ahora cada celda pasa por
    _excel_valor_seguro, que antepone una comilla simple a esos valores.

    Usa RequestFactory + llamado directo a la vista (no self.client) por el
    mismo bug de compatibilidad Django 4.2 + Python 3.14 documentado en
    test_cliente_preseleccionado_xss.py.
    """

    PAYLOAD = '=HYPERLINK("http://evil.com","click")'

    def setUp(self):
        self.factory = RequestFactory()

        rol, _ = Rol.objects.get_or_create(
            nombre='ADMIN',
            defaults={'descripcion': 'Rol admin para tests', 'activo': True}
        )
        perm, _ = Permiso.objects.get_or_create(
            codigo='reporte.export',
            defaults={'descripcion': 'reporte.export', 'activo': True}
        )
        RolPermiso.objects.get_or_create(rol=rol, permiso=perm)

        self.user = User.objects.create_user(
            username='admin_excel_test',
            password='testpass123'  # pragma: allowlist secret
        )
        UsuarioProfile.objects.update_or_create(
            usuario=self.user,
            defaults={'rol': rol}
        )

        self.cliente = Cliente.objects.create(
            nombre=self.PAYLOAD,
            cedula='1234567890',
            celular='3001234567',
        )

    def _primera_columna_de_datos(self, response):
        wb = load_workbook(io.BytesIO(response.content))
        ws = wb.active
        return [row[0].value for row in ws.iter_rows(min_row=2, max_row=ws.max_row)]

    def test_exportar_clientes_neutraliza_formula(self):
        request = self.factory.get('/exportar/clientes/')
        request.user = self.user
        response = exportar_clientes_excel(request)

        valores = self._primera_columna_de_datos(response)
        self.assertIn("'" + self.PAYLOAD, valores)
        self.assertNotIn(self.PAYLOAD, valores)

    def test_exportar_prestamos_neutraliza_formula(self):
        Prestamo.objects.create(
            cliente=self.cliente,
            monto_total=100000,
            interes_porcentaje=10,
            fecha_inicio=date.today(),
            fecha_fin_estimada=date.today().replace(year=date.today().year + 1),
            estado='ACTIVO',
        )
        request = self.factory.get('/exportar/prestamos/')
        request.user = self.user
        response = exportar_prestamos_excel(request)

        valores = self._primera_columna_de_datos(response)
        self.assertIn("'" + self.PAYLOAD, valores)
        self.assertNotIn(self.PAYLOAD, valores)
